from flask import Flask, request, send_file, jsonify
from flask_cors import CORS
import ezdxf
import math
import os
import re
from pyproj import Transformer
from docx import Document
from fpdf import FPDF
from PIL import Image

app = Flask(__name__)
CORS(app)

# --- CONFIGURAÇÃO GEOGRÁFICA SIRGAS 2000 ---
transformer = Transformer.from_crs("epsg:4674", "epsg:31984", always_xy=True)

def gms_para_decimal(gms_str, sentido):
    try:
        partes = re.findall(r"(\d+[\.,]?\d*)", gms_str)
        graus, minutos, segundos = float(partes[0]), float(partes[1]), float(partes[2].replace(',', '.'))
        decimal = graus + (minutos / 60) + (segundos / 3600)
        return -decimal if sentido in ['S', 'W', 'O'] else decimal
    except: return None

# --- ROTA 1: GEOCONVERTER ---
@app.route('/convert', methods=['POST'])
def convert():
    file = request.files['dxf_file']
    input_path, output_path = "temp_in.dxf", "temp_out.dxf"
    file.save(input_path)
    try:
        doc_in = ezdxf.readfile(input_path)
        msp_in = doc_in.modelspace()
        doc_out = ezdxf.new('R2010')
        blk = doc_out.blocks.new(name='SIG_PONTO')
        blk.add_circle((0, 0), radius=0.2)
        blk.add_attdef(tag='ID', text='ID', insert=(0.4, 0.4), dxfattribs={'height': 0.3})
        msp_out = doc_out.modelspace()
        contador_p = 1
        textos = msp_in.query('TEXT')
        for p in msp_in.query('POINT'):
            px, py = p.dxf.location.x, p.dxf.location.y
            id_f = next((t.dxf.text for t in textos if math.sqrt((px-t.dxf.insert.x)**2 + (py-t.dxf.insert.y)**2) < 3.0), f"P{contador_p}")
            if id_f == f"P{contador_p}": contador_p += 1
            br = msp_out.add_blockref('SIG_PONTO', (px, py), dxfattribs={'layer': 'VERTICES_ENGETECH'})
            br.add_auto_attribs({'ID': id_f})
        doc_out.saveas(output_path)
        return send_file(output_path, as_attachment=True, download_name="Engetech_Convertido.dxf")
    except Exception as e: return str(e), 500
    finally:
        if os.path.exists(input_path): os.remove(input_path)

# --- ROTA 2: GEOEXTRATOR (DOCX) ---
@app.route('/extract_docx', methods=['POST'])
def extract_docx():
    file = request.files['docx_file']
    docx_path = "temp_memorial.docx"
    file.save(docx_path)
    try:
        doc_word = Document(docx_path)
        texto = "\n".join([p.text for p in doc_word.paragraphs])
        lats = re.findall(r"(\d+°\d+'\d+[\.,]?\d*\"S)", texto)
        lons = re.findall(r"(\d+°\d+'\d+[\.,]?\d*\"W)", texto)
        vertices = []
        for i in range(min(len(lats), len(lons))):
            x, y = transformer.transform(gms_para_decimal(lons[i], 'W'), gms_para_decimal(lats[i], 'S'))
            vertices.append((x, y))
        doc_dxf = ezdxf.new('R2010')
        doc_dxf.layers.new('POLIGONAL', dxfattribs={'color': 1})
        msp = doc_dxf.modelspace()
        msp.add_lwpolyline(vertices, dxfattribs={'closed': True, 'layer': 'POLIGONAL'})
        output_dxf = "Engetech_Extraido.dxf"
        doc_dxf.saveas(output_dxf)
        return send_file(output_dxf, as_attachment=True)
    except Exception as e: return str(e), 500
    finally:
        if os.path.exists(docx_path): os.remove(docx_path)

# --- ROTA 3: GEOMEMORIAL (PDF + MARCA D'ÁGUA) ---
@app.route('/generate_pdf', methods=['POST'])
def generate_pdf():
    data = request.form
    logo = request.files.get('logo_file')
    pdf = FPDF()
    pdf.add_page()
    
    if logo:
        logo.save("logo_tmp.png")
        pdf.image("logo_tmp.png", x=50, y=100, w=110) # Marca d'água central
        
    pdf.set_font("Times", 'B', 14)
    pdf.cell(190, 10, "MEMORIAL DESCRITIVO", ln=True, align='C')
    pdf.set_font("Times", size=10)
    pdf.multi_cell(190, 6, f"PROPRIETÁRIO: {data.get('prop_nome')}\nCPF: {data.get('prop_cpf')}\nAREA: {data.get('area')} m²")
    pdf.ln(5)
    pdf.set_font("Times", size=11)
    pdf.multi_cell(190, 7, data.get('texto_descritivo'), align='J')
    
    path_pdf = "Memorial_Engetech.pdf"
    pdf.output(path_pdf)
    return send_file(path_pdf, as_attachment=True)

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=10000)
